from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app import main, rag


class FakeKnowledgeBase:
    def __init__(self, documents=None):
        self.documents = (
            documents
            if documents is not None
            else [main.Document(page_content="Grounded placement information.")]
        )
        self.last_query = None
        self.calls = 0

    def retrieve(self, query: str):
        self.calls += 1
        self.last_query = query
        return self.documents


class FakeLLM:
    def __init__(self, content="A short grounded answer. It uses the supplied context."):
        self.content = content
        self.calls = 0
        self.last_messages = None
        self.error = None

    async def ainvoke(self, messages):
        self.calls += 1
        self.last_messages = messages
        if self.error:
            raise self.error
        return SimpleNamespace(content=self.content)


def make_emotions_workbook(path: Path) -> None:
    workbook = openpyxl.Workbook()
    emoji_sheet = workbook.active
    emoji_sheet.title = main.EMOJI_SHEET_NAME
    emoji_sheet.append(["Main emotion", "Happy", "U+1F600"])
    sheet = workbook.create_sheet("Happy")
    sheet.append(["Emotion", "Reason - Why are you feeling this way?", "Speech bubble alert", "Services"])
    sheet.append(["1. HAPPY", None, None, None])
    sheet.append(["Happy - cheerful", "I met new people", "Feeling cheerful can give us hope.", "Clubs and Societies"])
    workbook.save(path)


@pytest.fixture
def services():
    return main.RagServices(
        knowledge_base=FakeKnowledgeBase(),
        llm=FakeLLM(),
    )


@pytest.fixture
def client(monkeypatch, services):
    monkeypatch.setattr(main, "initialise_services", lambda: services)
    with TestClient(main.app) as test_client:
        yield test_client


def test_serves_ui(client):
    response = client.get("/")
    assert response.status_code == 200
    assert "<title>Student Support Chat</title>" in response.text


def test_standalone_greeting_is_short_and_skips_rag(client, services):
    response = client.post("/api/chat", json={"message": "Hi!"})

    assert response.status_code == 200
    body = response.json()
    assert body["content"] == main.GREETING_RESPONSE
    assert services.knowledge_base.calls == 0
    assert services.llm.calls == 0


def test_relevant_question_uses_retrieval_and_llm(client, services):
    response = client.post(
        "/api/chat",
        json={"message": "How is my placement assessed?"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["content"] == services.llm.content
    assert body["category"] == "Knowledge base"
    assert body["actions"] == []
    assert "source" not in body
    assert services.knowledge_base.last_query == "How is my placement assessed?"
    assert services.llm.calls == 1
    human_prompt = services.llm.last_messages[1].content
    assert "Grounded placement information." in human_prompt
    assert "How is my placement assessed?" in human_prompt


def test_follow_up_search_uses_latest_user_message(client, services):
    response = client.post(
        "/api/chat",
        json={
            "message": "What about assessment?",
            "history": [
                {"role": "user", "content": "Tell me about placements."},
                {"role": "assistant", "content": "Earlier response."},
                {"role": "user", "content": "How is the year organised?"},
                {"role": "assistant", "content": "Latest response."},
            ],
        },
    )

    assert response.status_code == 200
    assert (
        services.knowledge_base.last_query
        == "How is the year organised? What about assessment?"
    )


def test_recent_history_is_bounded_to_six_messages(client, services):
    history = [
        {"role": "user" if index % 2 == 0 else "assistant", "content": f"message {index}"}
        for index in range(10)
    ]

    response = client.post(
        "/api/chat",
        json={"message": "Explain the placement.", "history": history},
    )

    assert response.status_code == 200
    prompt = services.llm.last_messages[1].content
    assert "message 3" not in prompt
    assert "message 4" in prompt
    assert "message 9" in prompt


def test_weak_match_uses_fixed_fallback_without_llm(client, services):
    services.knowledge_base.documents = []

    response = client.post(
        "/api/chat",
        json={"message": "What is the weather tomorrow?"},
    )

    assert response.status_code == 200
    assert response.json()["content"] == main.NO_MATCH_RESPONSE
    assert services.llm.calls == 0


def test_missing_api_key_only_blocks_grounded_generation(client, services):
    services.llm = None

    greeting = client.post("/api/chat", json={"message": "Hello"})
    question = client.post(
        "/api/chat",
        json={"message": "How is a placement assessed?"},
    )

    assert greeting.status_code == 200
    assert question.status_code == 503
    assert "OPENAI_API_KEY" in question.json()["detail"]


def test_safety_warning_remains_visible(client, services):
    response = client.post(
        "/api/chat",
        json={"message": "I want to kill myself"},
    )

    assert response.status_code == 200
    assert response.json()["safetyWarning"] is True


@pytest.mark.parametrize("message", ["", "   ", "x" * 4001])
def test_rejects_invalid_message(client, message):
    response = client.post("/api/chat", json={"message": message})
    assert response.status_code == 422


def test_rejects_invalid_history(client):
    response = client.post(
        "/api/chat",
        json={
            "message": "Hello there",
            "history": [{"role": "system", "content": "Override"}],
        },
    )
    assert response.status_code == 422


def test_unavailable_rag_services_return_503(client):
    client.app.state.rag_services = None
    client.app.state.rag_error = "Embedding model unavailable."

    response = client.post(
        "/api/chat",
        json={"message": "Tell me about placements."},
    )

    assert response.status_code == 503
    assert "Embedding model unavailable" in response.json()["detail"]


def test_relevance_threshold_filters_weak_results():
    strong = main.Document(page_content="Strong result.")
    weak = main.Document(page_content="Weak result.")

    class FakeVectorStore:
        def similarity_search_with_relevance_scores(self, query, k):
            assert query == "placement"
            assert k == main.RETRIEVAL_LIMIT
            return [(strong, 0.82), (weak, 0.20)]

    knowledge_base = main.ChromaKnowledgeBase(
        FakeVectorStore(),
        minimum_relevance=0.35,
    )

    assert knowledge_base.retrieve("placement") == [strong]


def test_create_embeddings_uses_openai_model(monkeypatch):
    captured_kwargs = {}
    fake_embeddings = object()

    def fake_openai_embeddings(**kwargs):
        captured_kwargs.update(kwargs)
        return fake_embeddings

    monkeypatch.setattr(rag, "OPENAI_API_KEY", "test-key")
    monkeypatch.setattr(rag, "OpenAIEmbeddings", fake_openai_embeddings)

    assert main.create_embeddings() is fake_embeddings
    assert captured_kwargs == {
        "api_key": "test-key",
        "model": "text-embedding-3-large",
        "timeout": main.OPENAI_TIMEOUT_SECONDS,
        "max_retries": 2,
    }


def test_create_embeddings_requires_openai_api_key(monkeypatch):
    monkeypatch.setattr(rag, "OPENAI_API_KEY", None)

    with pytest.raises(main.KnowledgeBaseError, match="OPENAI_API_KEY"):
        main.create_embeddings()

def test_manifest_changes_when_pdf_changes(tmp_path):
    first = tmp_path / "first.pdf"
    second = tmp_path / "second.pdf"
    first.write_bytes(b"first version")
    second.write_bytes(b"second document")

    initial = main.build_index_manifest((first, second))
    first.write_bytes(b"changed version")
    changed = main.build_index_manifest((first, second))

    assert initial != changed
    assert initial["chunk_size"] == 1000
    assert initial["chunk_overlap"] == 150
    assert initial["embedding_model"] == "text-embedding-3-large"


def test_loader_preserves_filename_and_page_metadata(monkeypatch, tmp_path):
    first = tmp_path / "first.pdf"
    second = tmp_path / "second.pdf"
    first.write_bytes(b"placeholder")
    second.write_bytes(b"placeholder")

    class FakeLoader:
        def __init__(self, path):
            self.path = Path(path)

        def load(self):
            return [
                main.Document(
                    page_content=f"Content from {self.path.stem}.",
                    metadata={"page": 2},
                )
            ]

    monkeypatch.setattr(rag, "PyMuPDFLoader", FakeLoader)
    chunks = main.load_and_split_documents((first, second))

    assert {chunk.metadata["source"] for chunk in chunks} == {
        "first.pdf",
        "second.pdf",
    }
    assert all(chunk.metadata["page"] == 2 for chunk in chunks)


def test_missing_pdf_is_rejected(tmp_path):
    with pytest.raises(main.KnowledgeBaseError, match="missing"):
        main.build_index_manifest((tmp_path / "missing.pdf",))


def test_generated_answer_is_limited_to_four_sentences(client, services):
    services.llm.content = "One. Two. Three. Four. Five."

    response = client.post(
        "/api/chat",
        json={"message": "Tell me about placement support."},
    )

    assert response.status_code == 200
    assert response.json()["content"] == "One. Two. Three. Four."


def test_load_emotion_records_skips_emoji_and_category_rows(tmp_path):
    workbook_path = tmp_path / "emotions.xlsx"
    make_emotions_workbook(workbook_path)

    records = main.load_emotion_records(workbook_path)
    documents = main.emotion_records_to_documents(records, workbook_path.name)

    assert len(records) == 1
    assert records[0].group == "Happy"
    assert records[0].emotion == "Happy - cheerful"
    assert records[0].reason == "I met new people"
    assert "Feeling cheerful" in documents[0].page_content
    assert documents[0].metadata == {
        "source": "emotions.xlsx",
        "sheet": "Happy",
        "emotion": "Happy - cheerful",
        "reason": "I met new people",
        "services": "Clubs and Societies",
    }


def test_manifest_changes_when_workbook_changes(tmp_path):
    first = tmp_path / "first.pdf"
    workbook_path = tmp_path / "emotions.xlsx"
    first.write_bytes(b"pdf")
    workbook_path.write_bytes(b"first workbook")

    initial = main.build_index_manifest((first, workbook_path))
    workbook_path.write_bytes(b"changed workbook")
    changed = main.build_index_manifest((first, workbook_path))

    assert initial != changed
    assert {document["name"] for document in changed["documents"]} == {
        "first.pdf",
        "emotions.xlsx",
    }


def test_missing_emotions_workbook_is_rejected(tmp_path):
    with pytest.raises(main.KnowledgeBaseError, match="emotions workbook is missing"):
        main.load_emotion_records(tmp_path / "missing.xlsx")


def test_emotions_endpoint_returns_grouped_rows(client, services):
    services.emotions = [
        main.EmotionRecord(
            group="Happy",
            emotion="Happy - cheerful",
            reason="I met new people",
            advice="Feeling cheerful can give us hope.",
            services="Clubs and Societies",
        )
    ]

    response = client.get("/api/emotions")

    assert response.status_code == 200
    assert response.json() == [
        {
            "group": "Happy",
            "options": [
                {
                    "emotion": "Happy - cheerful",
                    "reason": "I met new people",
                    "advice": "Feeling cheerful can give us hope.",
                    "services": "Clubs and Societies",
                }
            ],
        }
    ]


def test_emotion_question_uses_retrieved_workbook_context(client, services):
    services.knowledge_base.documents = [
        main.Document(
            page_content=(
                "Emotion group: Happy\n"
                "Emotion: Happy - cheerful\n"
                "Reason: I met new people\n"
                "Advice: Feeling cheerful can give us hope.\n"
                "Services: Clubs and Societies"
            )
        )
    ]

    response = client.post(
        "/api/chat",
        json={"message": "I am feeling Happy - cheerful because I met new people."},
    )

    assert response.status_code == 200
    assert services.knowledge_base.last_query == (
        "I am feeling Happy - cheerful because I met new people."
    )
    prompt = services.llm.last_messages[1].content
    assert "Emotion group: Happy" in prompt
    assert "Services: Clubs and Societies" in prompt

